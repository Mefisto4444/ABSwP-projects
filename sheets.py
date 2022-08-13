import openpyxl
wb = openpyxl.Workbook()
sheet = wb[wb.sheetnames[0]]
for i in range(1,11):
    sheet[f"A{str(i)}"].value = i

from openpyxl.chart import BarChart, Reference, Series
refObj = Reference(sheet,min_col=1,min_row=1,max_col=1,max_row=10)

SeriesObj = Series(refObj, title="Test")

chartObj = BarChart()
chartObj.title = "test"
chartObj.append(SeriesObj)

sheet.add_chart(chartObj,'C5')

try:
    wb.save("creator.xlsx")
except PermissionError:
    print("Wyłącz podgląd pliku.")